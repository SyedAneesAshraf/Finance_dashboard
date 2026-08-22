"""
Build the Phase 7 Excel companion: a budget-vs-actual workbook over Gold-layer spend.

Design, per docs/00_design_decisions.md's pattern of stating assumptions rather than
hiding them:

  - The `Invoices` and `GL_Accounts` sheets are the Gold export data, imported verbatim
    (no pre-aggregation in Python) — this is exports/gold/fact_invoice.csv and
    exports/gold/dim_gl_account.csv, unmodified.
  - `Invoices.DEPARTMENT` / `COST_CENTER` are live VLOOKUP formulas against `GL_Accounts`,
    not values copied over — so re-running export_gold_tables.py and re-opening this
    workbook reflects the new data without touching a single formula. (VLOOKUP rather
    than XLOOKUP: XLOOKUP needs Microsoft 365 / Excel 2021+, and this file should open
    cleanly in whatever Excel a reviewer has.)
  - `Budget_vs_Actual.Actual` is a live SUMIF over the `Invoices` sheet. Nothing here is
    a hardcoded number pulled from Snowflake — Excel computes it itself.
  - `Budget_vs_Actual.Budget` is the one manually-entered column in this workbook, because
    a budget is a management target, not a derived figure. Values are a documented
    assumption (see BUDGET_ASSUMPTIONS below): a round 24-month departmental spend plan,
    picked to be in the same order of magnitude as actuals rather than a number pulled to
    make every row balance.

Run:
    .venv\\Scripts\\python.exe scripts\\build_excel_companion.py
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO_ROOT = Path(__file__).resolve().parent.parent
GOLD_DIR = REPO_ROOT / "exports" / "gold"
OUT_PATH = REPO_ROOT / "excel" / "FinanceProcurement_BudgetVsActual.xlsx"

# Documented assumption (Phase 7 exit test requires a "reasonable, documented" budget,
# not manual entry disguised as one). Set as a round 24-month departmental spend plan,
# independent of the actuals below, so real variance — over in some departments, under
# in others — falls out rather than being reverse-engineered to look tidy.
BUDGET_ASSUMPTIONS = {
    "Manufacturing": 2_900_000,
    "IT": 2_250_000,
    "HR": 600_000,
    "Corporate": 650_000,
    "Finance": 400_000,
    "Logistics": 175_000,
    "Facilities": 12_000,
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, size=9, color="595959")
CURRENCY_FMT = '$#,##0;[RED]-$#,##0'
PCT_FMT = '0.0%;[RED]-0.0%'


def coerce(value: str):
    """csv.reader hands back plain strings for everything, including amounts and keys.
    Written as-is, openpyxl stores them as text cells, and Excel does not auto-detect
    numbers on file load the way it does on manual entry — SUMIF/VLOOKUP over a column
    of text-that-looks-numeric silently treats it as non-numeric. Coerce anything
    numeric-looking so downstream formulas see real numbers."""
    if value == "":
        return None
    try:
        return int(value)
    except ValueError:
        pass
    try:
        return float(value)
    except ValueError:
        return value


def read_csv(name: str) -> tuple[list[str], list[list[object]]]:
    path = GOLD_DIR / name
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = [[coerce(v) for v in row] for row in reader]
    return header, rows


def write_raw_sheet(wb: Workbook, sheet_name: str, header: list[str], rows: list[list[str]]) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.append(header)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    for col_idx, col_name in enumerate(header, start=1):
        letter = get_column_letter(col_idx)
        width = max(12, min(22, len(col_name) + 2))
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def build_invoices_sheet(wb: Workbook, gl_header: list[str]) -> tuple[str, int, str, str]:
    header, rows = read_csv("fact_invoice.csv")
    write_raw_sheet(wb, "Invoices", header, rows)
    ws = wb["Invoices"]
    n = len(rows)
    last_row = n + 1

    account_key_col = get_column_letter(header.index("ACCOUNT_KEY") + 1)
    amount_col = get_column_letter(header.index("AMOUNT_USD") + 1)
    for r in range(2, last_row + 1):
        ws.cell(row=r, column=header.index("AMOUNT_USD") + 1).number_format = CURRENCY_FMT
        ws.cell(row=r, column=header.index("AMOUNT") + 1).number_format = '#,##0.00'

    # VLOOKUP rather than XLOOKUP: XLOOKUP requires Microsoft 365 / Excel 2021+, and this
    # workbook needs to open cleanly in whatever Excel a reviewer happens to have.
    gl_key_idx = gl_header.index("ACCOUNT_KEY")
    gl_last_col = get_column_letter(len(gl_header))
    gl_dept_rel = gl_header.index("DEPARTMENT") - gl_key_idx + 1
    gl_cc_rel = gl_header.index("COST_CENTER") - gl_key_idx + 1
    if gl_dept_rel < 1 or gl_cc_rel < 1:
        raise ValueError("DEPARTMENT/COST_CENTER must be at or after ACCOUNT_KEY for VLOOKUP")
    gl_key_col = get_column_letter(gl_key_idx + 1)

    dept_col_idx = len(header) + 1
    cc_col_idx = len(header) + 2
    dept_letter = get_column_letter(dept_col_idx)
    cc_letter = get_column_letter(cc_col_idx)

    ws.cell(row=1, column=dept_col_idx, value="DEPARTMENT").fill = HEADER_FILL
    ws.cell(row=1, column=dept_col_idx).font = HEADER_FONT
    ws.cell(row=1, column=cc_col_idx, value="COST_CENTER").fill = HEADER_FILL
    ws.cell(row=1, column=cc_col_idx).font = HEADER_FONT

    gl_range = f"GL_Accounts!${gl_key_col}$2:${gl_last_col}$1048576"
    for r in range(2, last_row + 1):
        ws.cell(
            row=r, column=dept_col_idx,
            value=f"=VLOOKUP({account_key_col}{r},{gl_range},{gl_dept_rel},FALSE)",
        )
        ws.cell(
            row=r, column=cc_col_idx,
            value=f"=VLOOKUP({account_key_col}{r},{gl_range},{gl_cc_rel},FALSE)",
        )
    ws.column_dimensions[dept_letter].width = 16
    ws.column_dimensions[cc_letter].width = 14

    return amount_col, last_row, dept_letter, cc_letter


def build_gl_accounts_sheet(wb: Workbook) -> tuple[list[str], list[list[object]]]:
    header, rows = read_csv("dim_gl_account.csv")
    write_raw_sheet(wb, "GL_Accounts", header, rows)
    return header, rows


def build_cost_center_sheet(
    wb: Workbook,
    amount_col: str,
    last_row: int,
    dept_col: str,
    cc_col: str,
    gl_header: list[str],
    gl_rows: list[list[object]],
) -> None:
    """Actual spend by cost center — the granularity the roadmap's Phase 7 spec actually
    asks for ('actual spend by GL account/cost center'); the Budget_vs_Actual sheet rolls
    up to department because that's what the required chart is specified against. Every
    figure here is a live SUMIFS, so it also gives the exit test's 'spot-check 3-4 cost
    centers against Snowflake' something concrete to check without leaving the workbook."""
    dept_idx = gl_header.index("DEPARTMENT")
    cc_idx = gl_header.index("COST_CENTER")
    pairs = sorted({(r[dept_idx], r[cc_idx]) for r in gl_rows})

    ws = wb.create_sheet("CostCenter_Detail")
    ws["A1"] = "Actual Spend by Cost Center"
    ws["A1"].font = TITLE_FONT

    headers = ["Department", "Cost Center", "Actual"]
    header_row = 3
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    first_data_row = header_row + 1
    for i, (dept, cc) in enumerate(pairs):
        r = first_data_row + i
        ws.cell(row=r, column=1, value=dept)
        ws.cell(row=r, column=2, value=cc)
        ws.cell(
            row=r, column=3,
            value=f'=SUMIFS(Invoices!${amount_col}$2:${amount_col}${last_row},'
                  f'Invoices!${dept_col}$2:${dept_col}${last_row},$A{r},'
                  f'Invoices!${cc_col}$2:${cc_col}${last_row},$B{r})',
        ).number_format = CURRENCY_FMT

    last_data_row = first_data_row + len(pairs) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=f"=SUM(C{first_data_row}:C{last_data_row})").number_format = CURRENCY_FMT
    ws.cell(row=total_row, column=3).font = Font(bold=True)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    table_ref = f"A{header_row}:C{last_data_row}"
    tbl = Table(displayName="CostCenterDetail", ref=table_ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tbl)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{first_data_row}"


def build_budget_sheet(wb: Workbook, amount_col: str, last_row: int, dept_col: str) -> None:
    ws = wb.create_sheet("Budget_vs_Actual", 0)  # first tab

    ws["A1"] = "Budget vs Actual — by Department"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "24-month window (2024-09 → 2026-08). Budget is a documented planning "
        "assumption (see scripts/build_excel_companion.py); Actual and Variance are "
        "live formulas over the Invoices sheet, which is an unmodified Gold-layer export."
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 28
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")

    headers = ["Department", "Budget", "Actual", "Variance", "Variance %"]
    header_row = 4
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    depts = list(BUDGET_ASSUMPTIONS.keys())
    first_data_row = header_row + 1
    for i, dept in enumerate(depts):
        r = first_data_row + i
        ws.cell(row=r, column=1, value=dept)
        ws.cell(row=r, column=2, value=BUDGET_ASSUMPTIONS[dept]).number_format = CURRENCY_FMT
        ws.cell(
            row=r, column=3,
            value=f'=SUMIF(Invoices!${dept_col}$2:${dept_col}${last_row},$A{r},'
                  f'Invoices!${amount_col}$2:${amount_col}${last_row})',
        ).number_format = CURRENCY_FMT
        ws.cell(row=r, column=4, value=f"=C{r}-B{r}").number_format = CURRENCY_FMT
        ws.cell(row=r, column=5, value=f"=IFERROR(D{r}/B{r},0)").number_format = PCT_FMT

    last_data_row = first_data_row + len(depts) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=f"=SUM(B{first_data_row}:B{last_data_row})").number_format = CURRENCY_FMT
    ws.cell(row=total_row, column=3, value=f"=SUM(C{first_data_row}:C{last_data_row})").number_format = CURRENCY_FMT
    ws.cell(row=total_row, column=4, value=f"=C{total_row}-B{total_row}").number_format = CURRENCY_FMT
    ws.cell(row=total_row, column=5, value=f"=IFERROR(D{total_row}/B{total_row},0)").number_format = PCT_FMT
    for col in range(1, 6):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    # Conditional formatting: over budget = red, at/under budget = green.
    # Differential (CF) fills read color from end_color/bgColor, not fgColor as a normal
    # cell fill would — set both to the same value so the paint direction can't matter.
    variance_range = f"D{first_data_row}:D{last_data_row}"
    red_fill = PatternFill(fill_type="solid", start_color="F8CBAD", end_color="F8CBAD")
    green_fill = PatternFill(fill_type="solid", start_color="C6E0B4", end_color="C6E0B4")
    ws.conditional_formatting.add(
        variance_range,
        FormulaRule(formula=[f"D{first_data_row}>0"], fill=red_fill),
    )
    ws.conditional_formatting.add(
        variance_range,
        FormulaRule(formula=[f"D{first_data_row}<=0"], fill=green_fill),
    )

    ws.column_dimensions["A"].width = 16
    for col in "BCDE":
        ws.column_dimensions[col].width = 14

    table_ref = f"A{header_row}:E{last_data_row}"
    tbl = Table(displayName="BudgetVsActual", ref=table_ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True, showFirstColumn=False
    )
    ws.add_table(tbl)

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Budget vs Actual by Department"
    chart.y_axis.title = "USD"
    chart.x_axis.title = "Department"
    chart.style = 10

    data = Reference(ws, min_col=2, max_col=3, min_row=header_row, max_row=last_data_row)
    cats = Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 22
    chart.height = 11
    ws.add_chart(chart, f"G{header_row}")

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{first_data_row}"


def main() -> int:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    gl_header, gl_rows = build_gl_accounts_sheet(wb)
    amount_col, last_row, dept_col, cc_col = build_invoices_sheet(wb, gl_header)
    build_budget_sheet(wb, amount_col, last_row, dept_col)
    build_cost_center_sheet(wb, amount_col, last_row, dept_col, cc_col, gl_header, gl_rows)

    wb.active = 0
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
